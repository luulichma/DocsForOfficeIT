# -*- coding: utf-8 -*-
"""
Thư viện trợ giúp tạo định dạng DOCX chuẩn cho bộ tài liệu đào tạo Tin học văn phòng.
Tuân thủ chuẩn thể thức văn bản hành chính Việt Nam (Nghị định 30/2020/NĐ-CP) và phong cách trình bày giáo trình chuyên nghiệp.
"""

import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Bảng màu chuẩn chuyên nghiệp
COLOR_PRIMARY = RGBColor(13, 71, 161)     # Deep Blue
COLOR_SECONDARY = RGBColor(21, 101, 192)  # Medium Blue
COLOR_DARK = RGBColor(33, 33, 33)         # Near Black cho chữ
COLOR_MUTED = RGBColor(100, 116, 139)     # Xám ghi chú
COLOR_WARNING = RGBColor(185, 28, 28)     # Đỏ cảnh báo
COLOR_SUCCESS = RGBColor(21, 128, 61)     # Xanh lá

HEX_PRIMARY = "0D47A1"
HEX_BG_LIGHT = "F8FAFC"
HEX_BORDER = "CBD5E1"

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho một ô trong bảng."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt khoảng đệm (padding) bên trong ô (đơn vị dxa, 1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_border_left_only(cell, color_hex="0D47A1", sz="36"):
    """Tạo đường viền chỉ ở bên trái của ô (dùng cho Callout Box)."""
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_table_borders(table, color_hex="D1D5DB", sz="4"):
    """Đặt viền mỏng toàn bộ cho bảng."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def init_document():
    """Khởi tạo tài liệu với khổ giấy A4 và lề chuẩn Nghị định 30/2020/NĐ-CP."""
    doc = docx.Document()
    
    # Khổ giấy A4 và căn lề: Trên 20mm, Dưới 20mm, Trái 30mm, Phải 20mm
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(20)
        
        # Header / Footer
        section.different_first_page_header_footer = True
        
        # Đặt footer đánh số trang góc giữa
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Thêm field PAGE
        f_run = f_p.add_run("— ")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(11)
        f_run.font.color.rgb = COLOR_MUTED
        
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        f_p._p.append(fldSimple)
        
        f_run2 = f_p.add_run(" —")
        f_run2.font.name = "Times New Roman"
        f_run2.font.size = Pt(11)
        f_run2.font.color.rgb = COLOR_MUTED

    # Cài đặt phông chữ mặc định của tài liệu
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.3
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.space_before = Pt(0)
    
    return doc

def add_p(doc, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, color=None):
    """Thêm đoạn văn thông thường chuẩn hóa."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        if color:
            run.font.color.rgb = color
    return p

def add_heading_1(doc, text):
    """Tiêu đề Cấp 1 (Chương / Phần lớn)."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    """Tiêu đề Cấp 2 (Bài học / Đề mục lớn)."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    """Tiêu đề Cấp 3 (Mục nhỏ trong bài)."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_step(doc, step_num, step_title, details=None):
    """Định dạng một bước hướng dẫn thao tác chi tiết."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    r_badge = p.add_run(f"👉 Bước {step_num}: ")
    r_badge.bold = True
    r_badge.font.name = 'Times New Roman'
    r_badge.font.size = Pt(13)
    r_badge.font.color.rgb = COLOR_PRIMARY
    
    r_title = p.add_run(step_title)
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(13)
    
    if details:
        if isinstance(details, list):
            for d in details:
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.left_indent = Inches(0.45)
                p_sub.paragraph_format.space_before = Pt(1)
                p_sub.paragraph_format.space_after = Pt(2)
                r_dot = p_sub.add_run("• ")
                r_dot.bold = True
                r_dot.font.color.rgb = COLOR_SECONDARY
                r_content = p_sub.add_run(d)
                r_content.font.name = 'Times New Roman'
                r_content.font.size = Pt(12.5)
        else:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.45)
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(2)
            r_content = p_sub.add_run(details)
            r_content.font.name = 'Times New Roman'
            r_content.font.size = Pt(12.5)

def add_callout(doc, title, content_list, callout_type="note"):
    """
    Tạo hộp ghi chú chuyên nghiệp có viền màu bên trái và nền nhạt.
    callout_type:
      - 'note': Lưu ý chung (Màu xanh dương)
      - 'warning': Lỗi thường gặp & Cách khắc phục (Màu đỏ/cam)
      - 'tip': Mẹo thực chiến & Phím tắt nhanh (Màu xanh lá)
      - 'legal': Quy định pháp luật theo Nghị định 30/2020/NĐ-CP (Màu xanh đậm/tím)
    """
    type_configs = {
        "note": {"border": "1976D2", "fill": "F0F7FF", "title_color": RGBColor(25, 118, 210), "icon": "📌 "},
        "warning": {"border": "C62828", "fill": "FFF8F8", "title_color": RGBColor(198, 40, 40), "icon": "⚠️ "},
        "tip": {"border": "2E7D32", "fill": "F1F8E9", "title_color": RGBColor(46, 125, 50), "icon": "💡 "},
        "legal": {"border": "4527A0", "fill": "F3E5F5", "title_color": RGBColor(69, 39, 160), "icon": "⚖️ "}
    }
    cfg = type_configs.get(callout_type, type_configs["note"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Mm(160)
    set_cell_background(cell, cfg["fill"])
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)
    set_cell_border_left_only(cell, color_hex=cfg["border"], sz="36")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"{cfg['icon']}{title}")
    run_t.bold = True
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(13)
    run_t.font.color.rgb = cfg["title_color"]
    
    if isinstance(content_list, str):
        content_list = [content_list]
        
    for item in content_list:
        p_c = cell.add_paragraph()
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(2)
        p_c.paragraph_format.line_spacing = 1.25
        run_c = p_c.add_run(item)
        run_c.font.name = 'Times New Roman'
        run_c.font.size = Pt(12)
        run_c.font.color.rgb = COLOR_DARK
        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def add_image(doc, image_path, caption="", width=Mm(145)):
    """Chèn hình ảnh minh họa thực tế vào tài liệu DOCX với chú thích căn giữa."""
    import os
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run = p_img.add_run()
        run.add_picture(image_path, width=width)
        
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(0)
            p_cap.paragraph_format.space_after = Pt(6)
            r_cap = p_cap.add_run(f"Hình minh họa: {caption}")
            r_cap.italic = True
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(11)
            r_cap.font.color.rgb = COLOR_MUTED
    else:
        add_screenshot_placeholder(doc, caption or image_path)

def add_screenshot_placeholder(doc, description):
    """Thêm khung mô tả vị trí ảnh chụp màn hình cần bổ sung."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Mm(160)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    run_icon = p.add_run("📷 [VỊ TRÍ ẢNH CHỤP MÀN HÌNH MINH HỌA]\n")
    run_icon.bold = True
    run_icon.font.name = 'Times New Roman'
    run_icon.font.size = Pt(11)
    run_icon.font.color.rgb = RGBColor(71, 85, 105)
    
    run_desc = p.add_run(f"Mô tả chi tiết: {description}")
    run_desc.italic = True
    run_desc.font.name = 'Times New Roman'
    run_desc.font.size = Pt(11)
    run_desc.font.color.rgb = RGBColor(100, 116, 139)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def add_styled_table(doc, headers, rows, col_widths=None):
    """Tạo bảng dữ liệu chuyên nghiệp với tiêu đề nền xanh chữ trắng."""
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color_hex="CBD5E1", sz="4")
    
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
            
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if isinstance(val, (int, float)) or (isinstance(val, str) and val.isdigit()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11.5)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def add_checklist_table(doc, criteria_list):
    """Tạo bảng tiêu chí tự kiểm tra cuối bài học (Checklist)."""
    headers = ["STT", "Tiêu chí tự kiểm tra", "Kết quả mong đợi", "Tự đánh giá"]
    rows = []
    for idx, (crit, expected) in enumerate(criteria_list, 1):
        rows.append([str(idx), crit, expected, "[  ] Đạt   [  ] Chưa đạt"])
    
    col_widths = [Mm(15), Mm(65), Mm(55), Mm(25)]
    add_styled_table(doc, headers, rows, col_widths)

def add_cover_page(doc, volume_number, volume_title, subtitle, author="Nguyễn Thế Chiến"):
    """Tạo trang bìa mỹ thuật trang nhã, nghiêm túc."""
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        
    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_org = p_org.add_run("TÀI LIỆU HƯỚNG DẪN THỰC HÀNH CÔNG NGHỆ THÔNG TIN VĂN PHÒNG\n")
    r_org.bold = True
    r_org.font.name = 'Times New Roman'
    r_org.font.size = Pt(13)
    r_org.font.color.rgb = COLOR_SECONDARY
    
    r_line = p_org.add_run("———————————————\n\n\n")
    r_line.font.color.rgb = COLOR_MUTED
    
    p_vol = doc.add_paragraph()
    p_vol.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_vol = p_vol.add_run(f"TẬP {volume_number}\n")
    r_vol.bold = True
    r_vol.font.name = 'Times New Roman'
    r_vol.font.size = Pt(24)
    r_vol.font.color.rgb = COLOR_PRIMARY
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run(f"{volume_title.upper()}\n")
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(f"{subtitle}\n")
    r_sub.italic = True
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = COLOR_MUTED
    
    for _ in range(6):
        doc.add_paragraph()
        
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_author = p_footer.add_run(f"Người biên soạn: {author}\n")
    r_author.bold = True
    r_author.font.name = 'Times New Roman'
    r_author.font.size = Pt(13)
    r_author.font.color.rgb = COLOR_PRIMARY
    
    r_year = p_footer.add_run("Lưu hành nội bộ — Năm 2026")
    r_year.font.name = 'Times New Roman'
    r_year.font.size = Pt(11.5)
    r_year.font.color.rgb = COLOR_MUTED
    
    doc.add_page_break()
