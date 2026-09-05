# -*- coding: utf-8 -*-
"""
Module biên soạn PHỤ LỤC TẬP II:
- Phụ lục 1: Bảng phím tắt tiện ích Google Workspace
- Phụ lục 2: Quy tắc bảo mật và an toàn dữ liệu đám mây
"""

from docx_helper import (
    add_heading_1, add_heading_2, add_heading_3, add_p,
    add_styled_table, add_callout
)
from docx.shared import Mm

def build_vol2_appendix(doc):
    add_heading_1(doc, "PHỤ LỤC TẬP II: CẨM NANG CỘNG TÁC SỐ VÀ BẢO MẬT DỮ LIỆU ĐÁM MÂY")
    
    # PHỤ LỤC 1
    add_heading_2(doc, "Phụ lục 1: Bảng tổng hợp phím tắt quyền năng trên Google Docs và Google Sheets")
    add_p(doc, "Những phím tắt đặc thù trên trình duyệt web giúp bạn thao tác văn bản và bảng tính trực tuyến mượt mà như dùng phần mềm cài đặt trên máy tính.")
    
    headers_keys = ["Tổ hợp phím tắt", "Ứng dụng", "Tác dụng thực chiến"]
    rows_keys = [
        ["docs.new", "Trình duyệt Web", "Gõ vào thanh địa chỉ Chrome để mở ngay 1 tài liệu Google Docs mới."],
        ["sheets.new", "Trình duyệt Web", "Gõ vào thanh địa chỉ Chrome để mở ngay 1 bảng tính Google Sheets mới."],
        ["Ctrl + Alt + M", "Docs / Sheets", "TẠO BÌNH LUẬN (Comment) ngay tại vị trí con trỏ để góp ý."],
        ["Ctrl + Shift + J", "Google Docs", "Căn đều hai bên lề văn bản (Justify) trên web."],
        ["Ctrl + Shift + C", "Google Docs", "Đếm nhanh số từ và số trang trong tài liệu."],
        ["Ctrl + Alt + Shift + H", "Docs / Sheets", "MỞ LỊCH SỬ PHIÊN BẢN (Version History) để khôi phục bài cũ."],
        ["Ctrl + K", "Docs / Sheets", "Chèn đường liên kết (Hyperlink) vào văn bản."],
        ["Ctrl + /", "Docs / Sheets", "Mở bảng tra cứu toàn bộ danh mục phím tắt của Google."],
        ["Phím Tab", "Google Sheets", "Chấp nhận công thức gợi ý thông minh của Google."],
        ["Ctrl + Space", "Google Sheets", "Chọn nhanh toàn bộ một cột."],
        ["Shift + Space", "Google Sheets", "Chọn nhanh toàn bộ một hàng."]
    ]
    col_widths_keys = [Mm(45), Mm(32), Mm(83)]
    add_styled_table(doc, headers_keys, rows_keys, col_widths_keys)

    # PHỤ LỤC 2
    add_heading_2(doc, "Phụ lục 2: Quy tắc vàng bảo đảm an toàn và bảo mật dữ liệu công vụ trên đám mây")
    add_p(doc, "Làm việc trực tuyến mang lại sự tiện lợi vượt bậc nhưng cũng đòi hỏi ý thức bảo mật thông tin nghiêm ngặt để không làm lộ lọt dữ liệu của cơ quan, tổ chức.")

    add_callout(doc, "5 NGUYÊN TẮC BẢO MẬT DỮ LIỆU CÔNG VỤ DÀNH CHO CÁN BỘ VĂN PHÒNG", [
        "1. Luôn chia sẻ đích danh qua địa chỉ Email: Tránh tối đa việc bật chế độ 'Bất kỳ ai có đường liên kết đều xem được' đối với các tài liệu nội bộ hoặc có chứa số liệu tài chính, danh sách cá nhân. Khi gửi đích danh vào email của đồng nghiệp, chỉ người đó mới mở được tài liệu.",
        "2. Kiểm soát chặt chẽ quyền 'Người chỉnh sửa' (Editor): Chỉ cấp quyền Chỉnh sửa cho những người trực tiếp cùng soạn thảo. Đối với cán bộ phối hợp góp ý, hãy chỉ cấp quyền 'Người nhận xét' (Commenter) để tránh bị vô tình xóa mất dữ liệu gốc.",
        "3. Tắt tính năng Cho phép tải xuống và Sao chép: Đối với tài liệu mật hoặc tài liệu quan trọng: Trong cửa sổ Chia sẻ > Bấm vào biểu tượng bánh răng Cài đặt ở góc trên bên phải > BỎ TÍCH CHỌN ở ô 'Người xem và người nhận xét có thể nhìn thấy tùy chọn tải xuống, in và sao chép'.",
        "4. Thu hồi quyền truy cập khi có thay đổi nhân sự: Khi một cán bộ chuyển công tác hoặc chuyển sang phòng ban khác, người quản trị tài liệu cần vào danh sách Chia sẻ và xóa ngay email của cán bộ đó để bảo đảm an toàn dữ liệu nội bộ.",
        "5. Kích hoạt bảo mật 2 lớp (2-Step Verification) cho tài khoản Google: Bắt buộc cài đặt xác thực 2 bước qua điện thoại di động để dù kẻ gian có biết được mật khẩu hòm thư cũng không thể đăng nhập và đánh cắp tài liệu của bạn."
    ], callout_type="legal")
